from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO = Path(__file__).resolve().parents[2]
DOCS = REPO / "docs"
FIGS = REPO / "results" / "figures"

MAIN_OUT = DOCS / "main_report_final.docx"
SUPP_OUT = DOCS / "supplementary_materials.docx"


TITLE = "Benchmarking single-cell foundation model embeddings for fine-grained annotation of disease-associated NK/ILC states"

ABSTRACT = (
    "Frozen embeddings from scGPT, Geneformer and UCE were benchmarked against conventional classifiers "
    "for annotating NK/ILC states in juvenile idiopathic arthritis single-cell RNA-seq data. A dataset-specific "
    "Logistic Regression baseline remained strongest for fine-grained labels, whereas scGPT and UCE performed "
    "well with nonlinear probes and very strongly for broad NK/ILC labels. These results show that foundation "
    "models capture lineage structure but need task adaptation for subtle disease-associated states."
)

MAIN_SECTIONS = {
    "Background": [
        "Single-cell RNA sequencing can resolve immune heterogeneity at cellular resolution, but assigning cell states remains a major bottleneck. Manual annotation is slow and partly subjective, especially in inflammatory disease where transcriptional states may not align cleanly with canonical cell-type markers.",
        "Foundation models trained on large single-cell corpora, including scGPT [1], Geneformer [2] and Universal Cell Embedding (UCE) [3], aim to learn transferable cell representations. Their reported strengths include cell-type annotation, gene-network inference and transfer learning across datasets. However, their practical value depends on whether pretrained representations help on the target biological task, rather than only on broad benchmark labels.",
        "This project evaluates that question using NK and innate lymphoid cell (ILC) profiles from juvenile idiopathic arthritis (JIA), including both blood and synovial-fluid samples. Synovial fluid is biologically important because it represents the local inflammatory compartment of the affected joint, where immune cells may show activation, migration or tissue-resident programmes that are less apparent in blood. The benchmark is therefore harder than a generic PBMC annotation task: broad NK/ILC separation is biologically clear, whereas fine-grained labels such as C1, C2, C3, CD56bright NK, CD56dim NK, Cycling NK, ILC2, ILC3 and Naive-like ILC may reflect subtler subtype, disease-associated or tissue-context variation.",
    ],
    "Aims": [
        "The main aim was to test whether frozen single-cell foundation model embeddings improve automatic annotation of fine-grained NK/ILC states compared with simple, dataset-specific machine-learning baselines.",
        "Three foundation models were evaluated on the same train/test split: scGPT, Geneformer V1-10M and UCE. Their embeddings were kept frozen and classified with Logistic Regression or a multilayer perceptron (MLP). Traditional baselines used PCA-reduced normalized expression with Logistic Regression or k-nearest neighbours (kNN). Performance was assessed using accuracy, macro F1 and weighted F1. Additional analyses asked whether performance changed with label granularity and whether traditional baselines generalized from blood to synovial fluid.",
    ],
    "Results": [
        "After preprocessing, the benchmark dataset contained 16,836 cells and 22,144 genes. The stratified split used 13,468 training cells and 3,368 test cells. Fine-label support was uneven, ranging from 406 ILC3 cells to 3,000 cells for C1, C2 and CD56dim NK, so weighted F1 was used as the primary ranking metric and macro F1 was retained to check performance on smaller classes.",
        "Table 1 summarizes all results. For fine-grained celltype_l4 annotation, PCA plus Logistic Regression was strongest, with accuracy 0.9068, macro F1 0.8952 and weighted F1 0.9065. kNN on PCA features was slightly lower (weighted F1 0.8875). This matters because the baseline was a supervised model fitted directly to the study-specific label space, not a trivial comparator.",
        "Figure 1 summarizes the celltype_l4 weighted F1 comparison across the eight random-split models.",
        "Frozen foundation-model embeddings were informative but did not exceed the best PCA baseline on celltype_l4. With Logistic Regression probes, UCE performed best among the foundation models (weighted F1 0.8376), followed by scGPT (0.7759) and Geneformer V1-10M (0.7018). MLP probes improved scGPT to 0.8583 and UCE to 0.8558, while Geneformer changed little (0.7010), suggesting that nonlinear probing helped scGPT and UCE more than Geneformer.",
        "Label granularity had a strong effect. For broad celltype_l3 labels, scGPT and UCE Logistic Regression probes reached weighted F1 values of 0.9582 and 0.9741, respectively, much higher than their celltype_l4 results. Confusion matrices showed recurrent fine-state errors: C1 was often predicted as CD56bright NK, C1/C2 and C2/C3 were confused, and Naive-like ILC was sometimes predicted as C1 in foundation-model MLPs (Supplementary Table 2). CD56dim NK and ILC2 showed relatively strong per-class performance (Supplementary Figure 1; Supplementary Table 2).",
        "Per-class F1 for selected celltype_l4 models is provided in Supplementary Figure 1; the underlying verified confusion patterns are reported in Supplementary Table 2.",
        "Cross-tissue evaluation showed substantial domain shift. When trained on blood and tested on synovial fluid, Logistic Regression dropped from weighted F1 0.9065 in the random split to 0.7535, while kNN dropped from 0.8875 to 0.4376. Foundation-model cross-tissue transfer was not completed, so this remains a follow-up analysis rather than a conclusion from the present benchmark.",
    ],
    "Discussion": [
        "The results support a cautious interpretation of single-cell foundation models for this task. PCA plus Logistic Regression was the strongest celltype_l4 model in Table 1, but this is not simply a case of a small model beating larger models. It is a task-specific supervised classifier trained to reproduce the study's fine annotation scheme. Frozen scGPT, Geneformer and UCE embeddings instead represent expression profiles through pretrained spaces optimized for broader transfer, so they may not preserve every dataset-specific boundary needed for C1/C2/C3 or tissue-influenced NK states.",
        "The label-granularity results are the main biological finding. scGPT and UCE performed much better on broad celltype_l3 labels than on celltype_l4 (Table 1), indicating that frozen embeddings captured major NK versus ILC structure more reliably than fine subtype or state boundaries. The confusion patterns support this: errors clustered around C1, C2, C3, CD56bright NK and Naive-like ILC, while CD56dim NK and ILC2 were more consistently recovered (Supplementary Figure 1; Supplementary Table 2). These errors could reflect transcriptional overlap, tissue-context effects, or ambiguity in the original fine labels. In JIA, synovial fluid is the local inflammatory joint compartment, so immune cells may carry activation, migration or tissue-context signals that are less visible in blood [citation needed].",
        "The benchmark also has important fairness and statistical limits. The fairest direct comparison is Logistic Regression on PCA features versus Logistic Regression probes on frozen embeddings; MLP probes were not applied to the PCA baseline, and celltype_l3 was evaluated only for scGPT and UCE. The results are single-split point estimates without confidence intervals, so close differences such as scGPT-MLP versus UCE-MLP should not be overinterpreted. Future work should fine-tune scGPT or UCE, test foundation-model blood-to-synovial and cross-cohort transfer, quantify uncertainty, and validate whether the fine labels correspond to reproducible NK/ILC biology rather than study-specific annotation structure.",
    ],
}

TABLE_ROWS = [
    ["Logistic Regression baseline", "PCA normalized expression", "LR", "celltype_l4", "0.9068", "0.8952", "0.9065"],
    ["kNN baseline", "PCA normalized expression", "kNN", "celltype_l4", "0.8890", "0.8751", "0.8875"],
    ["scGPT", "Frozen embeddings, 512d", "LR", "celltype_l4", "0.7782", "0.7566", "0.7759"],
    ["scGPT", "Frozen embeddings, 512d", "MLP", "celltype_l4", "0.8584", "0.8381", "0.8583"],
    ["scGPT", "Frozen embeddings, 512d", "LR", "celltype_l3", "0.9561", "0.8998", "0.9582"],
    ["Geneformer V1-10M", "Frozen embeddings, 256d", "LR", "celltype_l4", "0.7007", "0.6511", "0.7018"],
    ["Geneformer V1-10M", "Frozen embeddings, 256d", "MLP", "celltype_l4", "0.7046", "0.6594", "0.7010"],
    ["UCE", "Frozen embeddings, 1280d", "LR", "celltype_l4", "0.8382", "0.8031", "0.8376"],
    ["UCE", "Frozen embeddings, 1280d", "MLP", "celltype_l4", "0.8566", "0.8445", "0.8558"],
    ["UCE", "Frozen embeddings, 1280d", "LR", "celltype_l3", "0.9733", "0.9364", "0.9741"],
    ["LR cross-tissue", "PCA normalized expression", "LR", "celltype_l4", "0.7504", "0.6706", "0.7535"],
    ["kNN cross-tissue", "PCA normalized expression", "kNN", "celltype_l4", "0.3965", "0.4446", "0.4376"],
]

REFERENCES = [
    "Cui, H., Wang, C., Maan, H. et al. scGPT: toward building a foundation model for single-cell multi-omics using generative AI. Nature Methods 21, 1470-1480 (2024). https://doi.org/10.1038/s41592-024-02201-0",
    "Theodoris, C.V., Xiao, L., Chopra, A. et al. Transfer learning enables predictions in network biology. Nature 618, 616-624 (2023). https://doi.org/10.1038/s41586-023-06139-9",
    "Rosen, Y., Roohani, Y., Agrawal, A. et al. Universal Cell Embeddings: a foundation model for cell biology. bioRxiv (2023). https://doi.org/10.1101/2023.11.28.568918",
    "Wolf, F.A., Angerer, P. and Theis, F.J. SCANPY: large-scale single-cell gene expression data analysis. Genome Biology 19, 15 (2018). https://doi.org/10.1186/s13059-017-1382-0",
    "Pedregosa, F., Varoquaux, G., Gramfort, A. et al. Scikit-learn: machine learning in Python. Journal of Machine Learning Research 12, 2825-2830 (2011).",
    "Paszke, A., Gross, S., Massa, F. et al. PyTorch: an imperative style, high-performance deep learning library. NeurIPS (2019).",
]

SUPP_METHODS = {
    "Data preprocessing": "The input was an h5ad JIA NK/ILC scRNA-seq dataset containing blood and synovial-fluid cells. Existing annotations were used for celltype_l3 and celltype_l4 labels. The final analyzed subset contained 16,836 cells and 22,144 genes, with median nCount_RNA 4,632, median nFeature_RNA 2,127 and median mitochondrial percentage 3.06. Normalized log1p expression was stored for classical models, and raw counts were retained where required by foundation-model workflows. PCA was computed on highly variable genes for baseline classifiers.",
    "Train/test split": "A fixed stratified 80/20 split was generated using celltype_l4 labels. The train and test sets contained 13,468 and 3,368 cells, respectively, and were reused for every model. Cross-tissue baseline experiments trained only on blood cells (n=6,762) and tested on synovial-fluid cells (n=10,074).",
    "Baseline classifiers": "Traditional models used 50 principal components from normalized expression. Logistic Regression used sklearn LogisticRegression with max_iter=5000. kNN used sklearn KNeighborsClassifier with n_neighbors=5. These baselines were evaluated on celltype_l4.",
    "Foundation model embeddings": "scGPT used the real human checkpoint with vocab.json, args.json and best_model.pt; 21,524 of 22,144 genes matched the model vocabulary and 512-dimensional cell embeddings were extracted. Geneformer used the real V1-10M model; Ensembl IDs were tokenized with Geneformer V1 dictionaries and 256-dimensional cell embeddings were extracted using EmbExtractor. UCE used the real 4-layer checkpoint with required token, chromosome, offset and protein-embedding files; gene identifiers were converted to symbols and 1,280-dimensional embeddings were read from obsm[\"X_uce\"]. No PCA fallback was used.",
    "Downstream evaluation": "Foundation-model weights were frozen. Logistic Regression served as a linear probe and an sklearn MLPClassifier with hidden layers of 256 and 128 units served as a nonlinear probe. Accuracy, macro F1 and weighted F1 were computed with sklearn metrics; per-class precision, recall and F1 were saved in classification-report tables. Weighted F1 was the primary ranking metric because class frequencies differed across labels. Code, environment files and output tables are provided in the local github_submission directory.",
}

REFLECTION = (
    "This project showed that strong baselines are essential when evaluating single-cell foundation models. The most useful result was not that a larger model wins, but that frozen embeddings excel at broad lineage recognition while fine subtype labels still favour supervised, dataset-specific learning. The main practical challenge was interoperability: scGPT, Geneformer and UCE required different gene identifiers, file formats, checkpoints and package versions. Course material on single-cell preprocessing, representation learning and benchmark design helped me keep the comparison fair and reproducible. Next, I would fine-tune scGPT or UCE, test foundation-model cross-tissue transfer from blood to synovial fluid, and validate whether C1/C2/C3 labels reflect stable biological programs."
)


def wc(text: str) -> int:
    return len(re.findall(r"\b[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)?\b", text))


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.1
    for name, size, color in [
        ("Title", 20, "0B2545"),
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(6)
    return doc


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(9)


def set_cell(cell, text: str, bold: bool = False, size: float = 7.2) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_main_table(doc: Document) -> None:
    add_caption(
        doc,
        "Table 1. Consolidated performance for traditional baselines and foundation-model probes. LR, Logistic Regression; MLP, multilayer perceptron. Cross-tissue rows were trained on blood and tested on synovial fluid. Rows from random-split celltype_l4, random-split celltype_l3 and cross-tissue celltype_l4 settings are different evaluation tasks.",
    )
    headers = ["Model", "Feature type", "Classifier", "Label", "Accuracy", "Macro F1", "Weighted F1"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, size=7.5)
        shade(table.rows[0].cells[i], "F2F4F7")
    for row in TABLE_ROWS:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell(cells[i], val)


def add_title(doc: Document, subtitle: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(TITLE)
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = q.add_run(subtitle)
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(85, 85, 85)


def build_main() -> None:
    doc = setup_doc()
    add_title(doc, "Main report")
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(ABSTRACT)
    doc.add_paragraph(f"Abstract word count: {wc(ABSTRACT)}")

    for section, paras in MAIN_SECTIONS.items():
        doc.add_heading(section, level=1)
        for para in paras:
            doc.add_paragraph(para)
            if para.startswith("Table 1 summarizes"):
                add_main_table(doc)
            if para.startswith("Figure 1 summarizes"):
                fig = FIGS / "figure1_celltype_l4_weighted_f1.png"
                doc.add_picture(str(fig), width=Inches(6.2))
                add_caption(
                    doc,
                    "Figure 1. Weighted F1 comparison across celltype_l4 models. Bars show weighted F1 for PCA-based baselines and foundation-model probes on the random-split celltype_l4 task. celltype_l3 and cross-tissue rows are excluded.",
                )

    main_wc = wc("\n".join(p for paras in MAIN_SECTIONS.values() for p in paras))
    doc.add_paragraph(f"Main text word count: {main_wc}")
    doc.add_heading("References", level=1)
    for i, ref in enumerate(REFERENCES, start=1):
        doc.add_paragraph(f"{i}. {ref}")
    doc.save(MAIN_OUT)


def build_supplementary() -> None:
    doc = setup_doc()
    add_title(doc, "Supporting materials")
    doc.add_heading("Supplementary Methods", level=1)
    for heading, text in SUPP_METHODS.items():
        doc.add_heading(heading, level=2)
        doc.add_paragraph(text)
    doc.add_paragraph(f"Supplementary methods word count: {wc(' '.join(SUPP_METHODS.values()))}")

    doc.add_heading("Reflection", level=1)
    doc.add_paragraph(REFLECTION)
    doc.add_paragraph(f"Reflection word count: {wc(REFLECTION)}")

    doc.add_heading("Supplementary Figures and Tables", level=1)
    doc.add_picture(str(FIGS / "figure2_per_class_f1_heatmap.png"), width=Inches(6.5))
    add_caption(
        doc,
        "Supplementary Figure 1. Per-class F1 for selected celltype_l4 models. Heatmap values show per-class F1 scores for the PCA + Logistic Regression baseline and the main MLP probes for scGPT, Geneformer V1-10M and UCE. Columns correspond to celltype_l4 labels.",
    )
    doc.add_picture(str(FIGS / "supplementary_table2_confusion_f1_summary.png"), width=Inches(6.5))
    add_caption(
        doc,
        "Supplementary Table 2. Verified confusion and per-class F1 summary for selected celltype_l4 models. Confusion entries show count/true-class support with percentages. Bidirectional C1/C2 and C2/C3 entries combine both directions. CD56dim NK and ILC2 values are per-class F1 scores from the corresponding classification reports.",
    )
    doc.save(SUPP_OUT)


def main() -> None:
    build_main()
    build_supplementary()
    print(MAIN_OUT)
    print(SUPP_OUT)


if __name__ == "__main__":
    main()
